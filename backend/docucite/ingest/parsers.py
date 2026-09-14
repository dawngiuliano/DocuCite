"""将 PDF、Word、Markdown 统一解析为 ParsedBlock。"""

from pathlib import Path
import re


from ..schemas import Document, Location, ParsedBlock, TableData


def parse_document(path: str | Path) -> tuple[Document, list[ParsedBlock]]:
    """解析一个文件，并返回“文档信息 + 原文块列表”。"""
    source = Path(path)
    # 根据扩展名选择解析器。解析器只在真正调用时导入对应依赖。
    kind = {".pdf": "pdf", ".docx": "docx", ".md": "md", ".xlsx": "xlsx"}.get(source.suffix.lower())
    if kind is None:
        raise ValueError(f"不支持的文件类型: {source.suffix}")
    document = Document(filename=source.name, file_type=kind)
    blocks = {"md": _parse_markdown, "docx": _parse_docx, "pdf": _parse_pdf, "xlsx": _parse_xlsx}[kind](source, document.doc_id)
    return document, blocks


def _parse_markdown(path: Path, doc_id: str) -> list[ParsedBlock]:
    """解析 Markdown：普通文字是文本块，Markdown 表格是表格块。"""
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks, headings, paragraph = [], [], []
    def flush():
        """把暂存的连续段落转换成一个文本块。"""
        if paragraph:
            text = "\n".join(paragraph).strip()
            if text:
                blocks.append(ParsedBlock(doc_id=doc_id, kind="text", text=text,
                    location=Location(heading_path=headings.copy(), paragraph_index=len(blocks)+1)))
            paragraph.clear()
    i = 0
    while i < len(lines):
        line = lines[i]
        match = re.match(r"^(#{1,6})\s+(.+?)\s*#*$", line)
        if match:
            flush(); level, title = len(match.group(1)), match.group(2).strip()
            headings[:] = headings[:level-1] + [title]; i += 1; continue
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+:?", lines[i+1]):
            flush(); header = _cells(line); rows, i = [], i + 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_cells(lines[i])); i += 1
            if rows:
                blocks.append(ParsedBlock(doc_id=doc_id, kind="table", table=TableData(header=header, rows=rows),
                    location=Location(heading_path=headings.copy(), table_index=len(blocks)+1)))
            continue
        if line.strip(): paragraph.append(line)
        else: flush()
        i += 1
    flush(); return blocks


def _cells(line: str) -> list[str]:
    """把一行 Markdown 表格拆成单元格，并去掉首尾空格。"""
    return [x.strip() for x in line.strip().strip("|").split("|")]


def _parse_docx(path: Path, doc_id: str) -> list[ParsedBlock]:
    """解析 Word 文档中的段落和表格。Word 通常没有可靠的页码信息。"""
    from docx import Document as WordDocument
    source = WordDocument(path); blocks = []; paragraph_no = 0
    # 按文档原有顺序遍历段落和表格，避免改变引用顺序。
    for item in source.element.body.iterchildren():
        if item.tag.endswith('}p'):
            text = ''.join(item.itertext()).strip(); paragraph_no += 1
            if text: blocks.append(ParsedBlock(doc_id=doc_id, kind="text", text=text, location=Location(paragraph_index=paragraph_no)))
        elif item.tag.endswith('}tbl'):
            table = next(t for t in source.tables if t._element == item)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows: blocks.append(ParsedBlock(doc_id=doc_id, kind="table", table=TableData(header=rows[0], rows=rows[1:] or [rows[0]]), location=Location(table_index=len(blocks)+1)))
    return blocks


def _parse_pdf(path: Path, doc_id: str) -> list[ParsedBlock]:
    """解析 PDF 的文字和表格，并使用真实页码作为位置。"""
    from pypdf import PdfReader
    import pdfplumber
    blocks = []
    reader = PdfReader(str(path))
    with pdfplumber.open(str(path)) as pdf:
        for page_no, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            if text: blocks.append(ParsedBlock(doc_id=doc_id, kind="text", text=text, location=Location(page=page_no)))
            for table_no, raw in enumerate(pdf.pages[page_no-1].extract_tables() or [], 1):
                rows = [[cell or "" for cell in row] for row in raw if row]
                if not rows: continue
                header, data = rows[0], rows[1:] or [rows[0]]
                blocks.append(ParsedBlock(doc_id=doc_id, kind="table", table=TableData(header=header, rows=data), location=Location(page=page_no, table_index=table_no)))
    return blocks


def _parse_xlsx(path: Path, doc_id: str) -> list[ParsedBlock]:
    """解析 Excel：每个非空工作表生成一个表格块。第一行是表头。"""
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks = []
    for table_index, sheet in enumerate(workbook.worksheets, 1):
        rows = [["" if v is None else str(v).strip() for v in row] for row in sheet.iter_rows(values_only=True)]
        while rows and not any(rows[-1]): rows.pop()
        if not rows: continue
        width = max(map(len, rows)); rows = [r + [""] * (width - len(r)) for r in rows]
        blocks.append(ParsedBlock(doc_id=doc_id, kind="table", table=TableData(header=rows[0], rows=rows[1:] or [rows[0]]), location=Location(heading_path=[sheet.title], table_index=table_index)))
    workbook.close()
    return blocks
